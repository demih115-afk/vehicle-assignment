# -*- coding: utf-8 -*-
"""엑셀 코스표 → routes.json 변환 스크립트"""
import openpyxl
import json
import re
import sys
import os

EXCEL_PATH = os.path.join(os.path.dirname(__file__), "차량 코스표 양식(260225).xlsx")
OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "data", "routes.json")

# 시트 인덱스 → 호차 정보 매핑 (이미지에서 확인한 정보)
VEHICLE_INFO = {
    0: {"number": 1, "area": "유곡동/우정동/태화동/센트리지", "driver": "박준성", "phone": "010-6677-9775"},
    1: {"number": 2, "area": "강변E편한/반구동/새치(학성동)", "driver": "서병진", "phone": "010-6566-7921"},
    2: {"number": 3, "area": "명촌동/학성초/반구동/남외동", "driver": "하수복", "phone": "010-4584-7036"},
    3: {"number": 5, "area": "복산동/약사동/래미안", "driver": "유종근", "phone": "010-6579-0441"},
    4: {"number": 6, "area": "성안동(성안초 방면)/센트리지", "driver": "김종철", "phone": "010-2850-0841"},
    5: {"number": 7, "area": "성안동(백양초 방면)/장현동", "driver": "박석칠", "phone": "010-6570-6243"},
    6: {"number": 8, "area": "서동/병영성/산전/약사아이파크", "driver": "김철현", "phone": "010-9332-2419"},
}


def format_time(val):
    """엑셀 시간값을 HH:MM 문자열로 변환"""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    # datetime.time 객체인 경우
    if hasattr(val, 'hour'):
        return f"{val.hour}:{val.minute:02d}"
    # "HH:MM:SS" 문자열인 경우
    m = re.match(r'(\d{1,2}):(\d{2})', s)
    if m:
        return f"{int(m.group(1))}:{m.group(2)}"
    return s


def format_schedule_key(val):
    """시간대 셀 값을 스케줄 키로 변환 (예: '07:00:00' → '7:00', '8:00\n(월.수.금)' → '8:00_월수금')"""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None

    is_saturday = False
    if '토요일' in s:
        is_saturday = True
        s = s.replace('토요일', '').strip()

    # 시간 추출
    time_match = re.search(r'(\d{1,2}):?(\d{2})?', s)
    if not time_match:
        return None
    hour = int(time_match.group(1))
    minute = time_match.group(2) or '00'
    time_str = f"{hour}:{minute}"

    # 요일 추출
    day_match = re.search(r'[(\uff08]([^\)\uff09]+)[)\uff09]', s)
    day_suffix = ""
    if day_match:
        days = day_match.group(1).replace('.', '').replace(',', '').replace(' ', '')
        day_suffix = f"_{days}"

    prefix = "토_" if is_saturday else ""
    return f"{prefix}{time_str}{day_suffix}"


def parse_sheet(ws):
    """시트 하나를 파싱하여 스케줄 딕셔너리 반환"""
    schedules = {}
    current_schedule_key = None
    max_col = ws.max_column or 23

    row_idx = 1
    max_row = ws.max_row or 100
    # 실제 데이터 범위 찾기 (빈 행이 많으므로 합리적 범위까지만)
    actual_max_row = min(max_row, 100)

    while row_idx <= actual_max_row:
        # B열(col 2)에서 시간대 확인
        b_cell = ws.cell(row=row_idx, column=2)
        c_cell = ws.cell(row=row_idx, column=3)

        b_val = b_cell.value
        c_val = str(c_cell.value).strip() if c_cell.value else ""

        # 새로운 시간대 시작 감지 (B열에 시간이 있는 경우)
        if b_val is not None and str(b_val).strip():
            new_key = format_schedule_key(b_val)
            if new_key:
                current_schedule_key = new_key
                if current_schedule_key not in schedules:
                    schedules[current_schedule_key] = []

        # 승차코스 행 감지
        if '승차코스' in c_val and current_schedule_key:
            stops_row = []
            for col in range(4, max_col + 1):
                cell = ws.cell(row=row_idx, column=col)
                name = cell.value
                if name is None or str(name).strip() == "":
                    continue
                name_str = str(name).strip().replace('\n', ' ')
                link = cell.hyperlink.target if cell.hyperlink else ""
                stops_row.append({"col": col, "name": name_str, "map_url": link})

            # 다음 행에서 승차시간 읽기
            next_row = row_idx + 1
            if next_row <= actual_max_row:
                next_c = ws.cell(row=next_row, column=3).value
                if next_c and '승차시간' in str(next_c):
                    times_by_col = {}
                    for col in range(4, max_col + 1):
                        t = ws.cell(row=next_row, column=col).value
                        ft = format_time(t)
                        if ft:
                            times_by_col[col] = ft

                    # 정류장과 시간 매칭
                    for stop in stops_row:
                        col = stop["col"]
                        stop["time"] = times_by_col.get(col, "")

                    row_idx += 1  # 시간 행 건너뛰기

            # 스케줄에 추가
            for stop in stops_row:
                schedules[current_schedule_key].append({
                    "stop": stop["name"],
                    "time": stop.get("time", ""),
                    "map_url": stop["map_url"]
                })

        row_idx += 1

    return schedules


def main():
    print(f"엑셀 파일 읽는 중: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)

    vehicles = []
    for idx, ws in enumerate(wb.worksheets):
        info = VEHICLE_INFO.get(idx)
        if not info:
            continue

        print(f"  {info['number']}호차 파싱 중...")
        schedules = parse_sheet(ws)

        # weekday / saturday 분리
        weekday = {}
        saturday = {}
        for key, stops in schedules.items():
            if not stops:
                continue
            if key.startswith("토_"):
                saturday[key.replace("토_", "")] = stops
            else:
                weekday[key] = stops

        vehicle = {
            "number": info["number"],
            "area": info["area"],
            "driver": info["driver"],
            "phone": info["phone"],
            "weekday": weekday,
            "saturday": saturday,
        }
        vehicles.append(vehicle)

        # 통계 출력
        total_stops = sum(len(s) for s in weekday.values()) + sum(len(s) for s in saturday.values())
        print(f"    평일 시간대: {list(weekday.keys())}")
        print(f"    토요일 시간대: {list(saturday.keys())}")
        print(f"    총 정류장 수: {total_stops}")

    result = {
        "academy_name": "더숲국어전문학원",
        "locations": {
            "본원": "더숲3관, 에스지, IBSI 영어, 엠플본관",
            "더숲1관": "신한은행 4층",
            "엠플2관": "피자스쿨",
            "더숲2관": "피자스쿨",
            "더숲국어": "알레르망 2층(래미안 맞은편)"
        },
        "vehicles": vehicles,
    }

    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n[OK] {OUTPUT_PATH} 생성 완료!")
    print(f"  차량 {len(vehicles)}대, 총 정류장 데이터 저장됨")


if __name__ == "__main__":
    main()
